"""Generate summary reports in XLSX and HTML formats.

Reads column definitions from ``config/report.yaml`` so the layout is
customisable without touching code.
"""

from __future__ import annotations

from datetime import datetime
import html
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from model_test_agent.state import ReportRow

_DEFAULT_REPORT_CFG = Path(__file__).resolve().parents[3] / "config" / "report.yaml"


class ReportGenerator:
    """Build XLSX and HTML reports from a list of :class:`ReportRow`.

    Parameters
    ----------
    config_path : Path | str
        Path to ``config/report.yaml``.
    output_dir : Path | str
        Directory where generated files are written.
    """

    def __init__(
        self,
        config_path: str | Path = _DEFAULT_REPORT_CFG,
        output_dir: str | Path = ".",
    ) -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        cfg_path = Path(config_path)
        if cfg_path.exists():
            with open(cfg_path, encoding="utf-8") as fh:
                self._cfg = yaml.safe_load(fh) or {}
        else:
            self._cfg = {}

    # ------------------------------------------------------------------
    # XLSX
    # ------------------------------------------------------------------

    def generate_xlsx(self, rows: list[ReportRow], filename: str | None = None) -> Path:
        """Write an XLSX report and return the file path."""
        xlsx_cfg = self._cfg.get("xlsx", {})
        xlsx_columns: list[dict[str, Any]] = xlsx_cfg.get("columns", self._default_columns())

        if filename is None:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            template = xlsx_cfg.get("filename_template", "test_report_{timestamp}.xlsx")
            filename = template.replace("{timestamp}", ts)

        wb = Workbook()
        ws = wb.active
        ws.title = xlsx_cfg.get("sheet_name", "Report")

        # --- Header row ---
        header_fill = PatternFill(
            start_color=xlsx_cfg.get("header_fill_color", "1F4E79"),
            end_color=xlsx_cfg.get("header_fill_color", "1F4E79"),
            fill_type="solid",
        )
        header_font = Font(
            bold=True,
            color=xlsx_cfg.get("header_font_color", "FFFFFF"),
            size=11,
        )
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_idx, col_def in enumerate(xlsx_columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_def["header"])
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col_def.get("width", 18)

        # --- Data rows ---
        alt_fill_color = xlsx_cfg.get("alternating_row_color", "D6E4F0")
        alt_fill = PatternFill(start_color=alt_fill_color, end_color=alt_fill_color, fill_type="solid")

        for row_idx, report_row in enumerate(rows, start=2):
            row_dict = self._row_to_dict(report_row)
            for col_idx, col_def in enumerate(xlsx_columns, start=1):
                value = row_dict.get(col_def["key"], "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

        # Freeze header
        ws.freeze_panes = "A2"

        out_path = self.output_dir / filename
        wb.save(str(out_path))
        return out_path

    # ------------------------------------------------------------------
    # HTML
    # ------------------------------------------------------------------

    def generate_html(
        self,
        rows: list[ReportRow],
        filename: str | None = None,
        summary: dict[str, int] | None = None,
        agent_info: dict[str, str] | None = None,
        source_info: dict[str, str] | None = None,
    ) -> Path:
        """Write a self-contained HTML report and return the file path."""
        html_cfg = self._cfg.get("html", {})
        xlsx_cfg = self._cfg.get("xlsx", {})
        xlsx_columns: list[dict[str, Any]] = xlsx_cfg.get("columns", self._default_columns())
        html_columns = [column for column in xlsx_columns if column.get("key") not in {"checked_by", "comment"}]

        if filename is None:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            template = html_cfg.get("filename_template", "test_report_{timestamp}.html")
            filename = template.replace("{timestamp}", ts)

        title = html_cfg.get("title", "Model Conversion Test Report")

        # Build table rows
        header_cells = "<th>Checked By</th><th>Comment</th>"
        log_inserted = False
        for column in html_columns:
            header_cells += f'<th>{html.escape(column["header"])}</th>'
            if column["key"] == "error_category":
                header_cells += "<th>Log</th>"
                log_inserted = True
        if not log_inserted:
            header_cells += "<th>Log</th>"
        body_rows = []
        log_cache: dict[str, str] = {}
        for i, report_row in enumerate(rows):
            rd = self._row_to_dict(report_row)
            categories = self._row_categories(report_row)
            row_classes = ["report-row"]
            if i % 2 == 1:
                row_classes.append("alt")
            row_id = f"row-{i}"
            log_id = f"log-{i}"
            checked_by_key = self._row_storage_key(report_row, i)
            checked_by = report_row.checked_by.strip()
            comment = report_row.comment.strip()
            cell_parts = [
                self._build_checked_by_cell(report_row, checked_by_key),
                self._build_comment_cell(report_row, checked_by_key),
            ]
            log_cell = f'<td class="log-actions">{self._build_log_actions(report_row, log_id)}</td>'
            log_inserted = False
            for column in html_columns:
                cell_parts.append(self._build_table_cell(column["key"], rd.get(column["key"], "")))
                if column["key"] == "error_category":
                    cell_parts.append(log_cell)
                    log_inserted = True
            if not log_inserted:
                cell_parts.append(log_cell)
            body_rows.append(
                f'  <tr id="{row_id}" class="{" ".join(row_classes)}"'
                f' data-status="{html.escape(report_row.status, quote=True)}"'
                f' data-categories="{html.escape("|".join(categories), quote=True)}"'
                f' data-checked="{str(bool(checked_by)).lower()}"'
                f' data-detail-target="{log_id if report_row.log_path else ""}">'
                f'{"".join(cell_parts)}</tr>'
            )
            if report_row.log_path:
                body_rows.append(
                    self._build_log_detail(report_row, len(html_columns) + 3, log_id, row_id, log_cache)
                )
        body_html = "\n".join(body_rows)

        # Statistics for the summary section
        summary = summary or {}
        total_models = summary.get("total_models", len({row.model_name for row in rows}))
        passed_models = summary.get("passed_models", 0)
        failed_models = summary.get("failed_models", len({row.model_name for row in rows}))
        agent_info = agent_info or {}
        source_info = source_info or {}
        by_category: dict[str, int] = {}
        by_status: dict[str, int] = {}
        for r in rows:
            by_category[r.error_category] = by_category.get(r.error_category, 0) + 1
            by_status[r.status] = by_status.get(r.status, 0) + 1

        stats_rows = "".join(
            "<tr class=\"filter-option\""
            f' data-filter-kind="category" data-filter-value="{html.escape(cat, quote=True)}">'
            f"<td>{html.escape(cat)}</td><td>{cnt}</td></tr>"
            for cat, cnt in sorted(by_category.items())
        )
        status_rows = "".join(
            "<tr class=\"filter-option\""
            f' data-filter-kind="status" data-filter-value="{html.escape(st, quote=True)}">'
            f"<td>{html.escape(st)}</td><td>{cnt}</td></tr>"
            for st, cnt in sorted(by_status.items())
        )
        html_output = _HTML_TEMPLATE.format(
            title=title,
            header_cells=header_cells,
            body_rows=body_html,
            total_models=total_models,
            passed_models=passed_models,
            failed_models=failed_models,
            stats_rows=stats_rows,
            status_rows=status_rows,
            classifier_model=html.escape(agent_info.get("classifier_model", "-")),
            debugger_model=html.escape(agent_info.get("debugger_model", "-")),
            assisted_fields=html.escape(agent_info.get("assisted_fields", "")),
            target_dir=html.escape(source_info.get("target_dir", "-")),
            layout_config=html.escape(source_info.get("layout_config", "built-in defaults")),
            discovery_rule=html.escape(source_info.get("discovery_rule", "")),
            source_tree=html.escape(source_info.get("source_tree", "")),
            timestamp=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            visible_rows=len(rows),
        )

        out_path = self.output_dir / filename
        out_path.write_text(html_output, encoding="utf-8")
        return out_path

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _row_to_dict(row: ReportRow) -> dict[str, Any]:
        return {
            "model_name": row.model_name,
            "log_path": row.log_path,
            "log_line": row.log_line,
            "package_summary": row.package_summary,
            "quantization": row.quantization,
            "has_test_data": row.has_test_data,
            "config_hints": row.config_hints,
            "error_category": row.error_category,
            "error_count": row.error_count,
            "key_log_snippet": row.key_log_snippet,
            "history_match": row.history_match,
            "suggested_fix": row.suggested_fix,
            "fix_executed": row.fix_executed,
            "fix_result": row.fix_result,
            "status": row.status,
            "checked_by": row.checked_by,
            "comment": row.comment,
        }

    @staticmethod
    def _default_columns() -> list[dict[str, Any]]:
        return [
            {"key": "model_name", "header": "Model Name", "width": 28},
            {"key": "package_summary", "header": "Package", "width": 28},
            {"key": "quantization", "header": "Quantization", "width": 14},
            {"key": "has_test_data", "header": "Has Test Data", "width": 13},
            {"key": "config_hints", "header": "Config Hints", "width": 24},
            {"key": "error_category", "header": "Error Category", "width": 18},
            {"key": "key_log_snippet", "header": "Key Log (truncated)", "width": 50},
            {"key": "history_match", "header": "Seen Before", "width": 12},
            {"key": "suggested_fix", "header": "Suggested Fix", "width": 72},
            {"key": "fix_executed", "header": "Fix Executed", "width": 12},
            {"key": "fix_result", "header": "Fix Result", "width": 18},
            {"key": "status", "header": "Status", "width": 12},
        ]

    @staticmethod
    def _build_table_cell(key: str, value: Any) -> str:
        text = html.escape(str(value))
        css_class = f' class="cell-{html.escape(key, quote=True)}"' if key else ""
        title = f' title="{text}"' if key in {"suggested_fix", "key_log_snippet"} and text else ""
        return f"<td{css_class}{title}>{text}</td>"

    @staticmethod
    def _build_log_actions(row: ReportRow, log_id: str) -> str:
        if not row.log_path:
            return '<span class="muted">-</span>'
        file_link = Path(row.log_path).resolve().as_uri()
        return (
            f'<button type="button" class="log-toggle" data-target="{log_id}">View Source</button>'
            f'<button type="button" class="log-link log-open"'
            f' data-file-link="{html.escape(file_link, quote=True)}"'
            f' data-log-path="{html.escape(row.log_path, quote=True)}">'
            "Open File</button>"
        )

    @staticmethod
    def _build_checked_by_cell(row: ReportRow, row_key: str) -> str:
        checked_by = row.checked_by.strip()
        checked_attr = " checked" if checked_by else ""
        label = html.escape(checked_by) if checked_by else "-"
        return (
            '<td class="checked-by-cell"'
            f' data-row-key="{html.escape(row_key, quote=True)}"'
            f' data-initial-checked-by="{html.escape(checked_by, quote=True)}">'
            '<label class="checked-by-toggle">'
            f'<input type="checkbox" class="checked-by-checkbox"{checked_attr}>'
            "<span>Checked</span>"
            "</label>"
            f'<div class="checked-by-name">{label}</div>'
            "</td>"
        )

    @staticmethod
    def _build_comment_cell(row: ReportRow, row_key: str) -> str:
        comment = html.escape(row.comment.strip())
        return (
            '<td class="comment-cell"'
            f' data-row-key="{html.escape(row_key, quote=True)}"'
            f' data-initial-comment="{comment}">'
            f'<textarea class="comment-input" rows="2" placeholder="Add comment">{comment}</textarea>'
            "</td>"
        )

    @classmethod
    def _build_log_detail(
        cls,
        row: ReportRow,
        colspan: int,
        log_id: str,
        parent_row_id: str,
        log_cache: dict[str, str],
    ) -> str:
        source = cls._read_log_source(row.log_path, log_cache)
        line_label = f"line {row.log_line}" if row.log_line else "unknown line"
        return (
            f'  <tr id="{log_id}" class="log-detail-row" data-parent-row="{parent_row_id}" hidden>'
            f'<td colspan="{colspan}">'
            f'<div class="log-detail-meta">{html.escape(row.log_path)} ({line_label})</div>'
            f'<pre class="log-detail"><code>{cls._format_log_source(source, row.log_line)}</code></pre>'
            "</td></tr>"
        )

    @staticmethod
    def _row_categories(row: ReportRow) -> list[str]:
        return [part.strip() for part in row.error_category.split(",") if part.strip()]

    @staticmethod
    def _row_storage_key(row: ReportRow, index: int) -> str:
        return "|".join([
            str(index),
            row.model_name,
            row.log_path,
            str(row.log_line),
            row.error_category,
        ])

    @staticmethod
    def _read_log_source(log_path: str, log_cache: dict[str, str]) -> str:
        if log_path in log_cache:
            return log_cache[log_path]
        path = Path(log_path)
        try:
            content = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            content = f"[Unable to read log file] {exc}"
        log_cache[log_path] = content
        return content

    @staticmethod
    def _format_log_source(source: str, highlight_line: int) -> str:
        if not source:
            return html.escape("[Empty log]")
        rendered_lines = []
        for index, line in enumerate(source.splitlines(), start=1):
            css_class = ' class="hit"' if index == highlight_line else ""
            rendered_lines.append(
                f'<span{css_class}><span class="ln">{index:>5}</span> {html.escape(line)}</span>'
            )
        return "\n".join(rendered_lines)


# ------------------------------------------------------------------
# HTML template (no external dependencies)
# ------------------------------------------------------------------

_HTML_TEMPLATE = """\
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<style>
  :root {{
    --primary: #1F4E79;
    --primary-light: #D6E4F0;
    --success: #27ae60;
    --danger: #e74c3c;
    --warning: #f39c12;
    --bg: #f5f6fa;
    --text: #2c3e50;
  }}
  * {{ margin: 0; padding: 0; box-sizing: border-box; }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    background: var(--bg);
    color: var(--text);
    padding: 24px;
  }}
  h1 {{
    color: var(--primary);
    margin-bottom: 8px;
    font-size: 1.6em;
  }}
  .meta {{ color: #7f8c8d; margin-bottom: 24px; font-size: 0.9em; }}
  .summary {{
    display: flex;
    gap: 24px;
    margin-bottom: 24px;
    flex-wrap: wrap;
  }}
  .summary-card {{
    background: #fff;
    border-radius: 8px;
    padding: 16px 24px;
    box-shadow: 0 2px 8px rgba(0,0,0,0.06);
    min-width: 200px;
  }}
  .summary-card.filter-card,
  .filter-option {{
    cursor: pointer;
    transition: transform 0.15s ease, box-shadow 0.15s ease, background 0.15s ease;
  }}
  .summary-card.filter-card:hover,
  .filter-option:hover {{
    transform: translateY(-1px);
  }}
  .summary-card.filter-card.active {{
    box-shadow: 0 0 0 2px rgba(31, 78, 121, 0.18), 0 8px 20px rgba(31, 78, 121, 0.12);
  }}
  .summary-card.agent {{
    min-width: 320px;
  }}
  .summary-card.source {{
    min-width: 360px;
  }}
  .summary-card h3 {{
    font-size: 0.85em;
    color: #7f8c8d;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-bottom: 8px;
  }}
  .summary-card .big {{ font-size: 2em; font-weight: 700; color: var(--primary); }}
  .agent-lines {{
    font-size: 0.92em;
    line-height: 1.6;
    color: #425466;
  }}
  .source-lines {{
    font-size: 0.92em;
    line-height: 1.6;
    color: #425466;
  }}
  .source-tree {{
    margin-top: 10px;
    padding: 12px;
    border-radius: 8px;
    background: #f6f8fb;
    color: #304050;
    font-size: 0.86em;
    line-height: 1.55;
    white-space: pre-wrap;
    font-family: ui-monospace, SFMono-Regular, Menlo, monospace;
  }}
  .agent-lines strong {{
    color: var(--primary);
  }}
  .source-lines strong {{
    color: var(--primary);
  }}
  .summary-card table {{ width: 100%; font-size: 0.9em; }}
  .summary-card td {{ padding: 2px 8px; }}
  .summary-card td:last-child {{ text-align: right; font-weight: 600; }}
  .filter-option.active {{
    background: rgba(31, 78, 121, 0.08);
  }}
  .toolbar {{
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 12px;
    flex-wrap: wrap;
    margin-bottom: 12px;
    color: #51606f;
    font-size: 0.92em;
  }}
  .toolbar-left {{
    display: flex;
    align-items: center;
    gap: 12px;
    flex-wrap: wrap;
  }}
  .filter-pills {{
    display: inline-flex;
    align-items: center;
    gap: 8px;
    flex-wrap: wrap;
  }}
  .filter-pills-label {{
    font-weight: 600;
    color: #425466;
  }}
  .toolbar button {{
    border: 0;
    border-radius: 999px;
    padding: 8px 12px;
    background: #e8f0f8;
    color: var(--primary);
    cursor: pointer;
    font-weight: 600;
  }}
  .table-wrapper {{
    overflow-x: auto;
    background: #fff;
    border-radius: 8px;
    box-shadow: 0 2px 8px rgba(0,0,0,0.06);
  }}
  table.report {{
    width: 100%;
    border-collapse: collapse;
    font-size: 0.85em;
  }}
  table.report th {{
    background: var(--primary);
    color: #fff;
    padding: 10px 12px;
    text-align: left;
    position: sticky;
    top: 0;
    white-space: nowrap;
    position: sticky;
  }}
  table.report th.resizable {{
    position: sticky;
  }}
  .column-resizer {{
    position: absolute;
    top: 0;
    right: -3px;
    width: 8px;
    height: 100%;
    cursor: col-resize;
    user-select: none;
    touch-action: none;
  }}
  table.report td {{
    padding: 8px 12px;
    border-bottom: 1px solid #ecf0f1;
    max-width: 400px;
    word-wrap: break-word;
    overflow-wrap: anywhere;
    vertical-align: top;
  }}
  table.report td.cell-suggested_fix {{
    min-width: 420px;
    max-width: 720px;
    white-space: pre-wrap;
    line-height: 1.55;
  }}
  table.report td.cell-fix_result,
  table.report td.cell-status {{
    white-space: nowrap;
    min-width: 150px;
  }}
  table.report td.cell-key_log_snippet {{
    max-width: 520px;
    white-space: pre-wrap;
  }}
  .checked-by-cell {{
    min-width: 170px;
    white-space: nowrap;
  }}
  .comment-cell {{
    min-width: 240px;
  }}
  .checked-by-toggle {{
    display: inline-flex;
    align-items: center;
    gap: 8px;
    font-weight: 600;
    color: var(--primary);
    cursor: pointer;
  }}
  .checked-by-checkbox {{
    width: 16px;
    height: 16px;
    accent-color: var(--primary);
    cursor: pointer;
  }}
  .checked-by-name {{
    margin-top: 6px;
    color: #51606f;
    font-size: 0.92em;
  }}
  .comment-input {{
    width: 100%;
    min-height: 54px;
    resize: vertical;
    border: 1px solid #d3dce6;
    border-radius: 8px;
    padding: 8px 10px;
    background: #fbfcfe;
    color: #2c3e50;
    font: inherit;
    line-height: 1.4;
  }}
  .comment-input:focus {{
    outline: none;
    border-color: rgba(31, 78, 121, 0.65);
    box-shadow: 0 0 0 3px rgba(31, 78, 121, 0.12);
    background: #fff;
  }}
  .log-actions {{
    white-space: nowrap;
    min-width: 180px;
  }}
  .log-toggle, .log-link {{
    display: inline-block;
    border-radius: 999px;
    padding: 6px 10px;
    font-size: 0.85em;
    text-decoration: none;
    margin-right: 8px;
  }}
  .log-toggle {{
    border: 0;
    cursor: pointer;
    background: var(--primary);
    color: #fff;
  }}
  .log-link {{
    border: 0;
    background: #eef4fb;
    color: var(--primary);
    cursor: pointer;
  }}
  .muted {{
    color: #95a5a6;
  }}
  .log-detail-row {{
    background: #fbfcfe;
  }}
  .log-detail-meta {{
    font-family: ui-monospace, SFMono-Regular, Menlo, monospace;
    color: #51606f;
    margin-bottom: 10px;
  }}
  .log-detail {{
    margin: 0;
    background: #0f1720;
    color: #e6edf3;
    padding: 14px;
    border-radius: 8px;
    overflow-x: auto;
    font-family: ui-monospace, SFMono-Regular, Menlo, monospace;
    line-height: 1.5;
  }}
  .log-detail code {{
    display: block;
  }}
  .log-detail code > span {{
    display: block;
    white-space: pre;
  }}
  .log-detail .ln {{
    color: #7c8a99;
    margin-right: 12px;
  }}
  .log-detail .hit {{
    background: rgba(41, 194, 166, 0.18);
  }}
  table.report tr.alt {{ background: var(--primary-light); }}
  table.report tr:hover {{ background: #eaf2f8; }}
  tr[hidden] {{ display: none !important; }}
  footer {{
    margin-top: 24px;
    text-align: center;
    color: #95a5a6;
    font-size: 0.8em;
  }}
  .toast {{
    position: fixed;
    right: 24px;
    bottom: 24px;
    max-width: 380px;
    padding: 12px 14px;
    border-radius: 10px;
    background: rgba(15, 23, 32, 0.94);
    color: #f4f8fb;
    box-shadow: 0 10px 24px rgba(0,0,0,0.22);
    font-size: 0.9em;
    line-height: 1.45;
    opacity: 0;
    transform: translateY(10px);
    pointer-events: none;
    transition: opacity 0.18s ease, transform 0.18s ease;
    z-index: 20;
  }}
  .toast.visible {{
    opacity: 1;
    transform: translateY(0);
  }}
</style>
<script>
  document.addEventListener("DOMContentLoaded", () => {{
    const filters = {{
      group: "all",
      category: "all",
      status: "all",
      checked: "all",
    }};
    const reportRows = Array.from(document.querySelectorAll("tr.report-row"));
    const detailRows = Array.from(document.querySelectorAll(".log-detail-row"));
    const checkedByCells = Array.from(document.querySelectorAll(".checked-by-cell"));
    const commentCells = Array.from(document.querySelectorAll(".comment-cell"));
    const filterState = document.getElementById("filter-state");
    const clearButton = document.getElementById("clear-filters");
    const toast = document.getElementById("toast");
    const usernameStorageKey = "model-test-agent:checked-by-user";
    const reportStorageKey = `model-test-agent:report-review:${{window.location.pathname}}`;
    const reviewApiUrl = `/api/review-state?report=${{encodeURIComponent(window.location.pathname)}}`;
    const reviewState = loadReviewState();
    const canOpenLocalFiles = window.location.protocol === "file:";
    let toastTimer = null;
    let saveTimer = null;

    function enableColumnResize() {{
      const table = document.querySelector("table.report");
      if (!table) return;
      const headers = Array.from(table.querySelectorAll("th"));
      headers.forEach((header) => {{
        header.classList.add("resizable");
        const handle = document.createElement("span");
        handle.className = "column-resizer";
        let startX = 0;
        let startWidth = 0;
        const onMove = (event) => {{
          const nextWidth = Math.max(90, startWidth + event.clientX - startX);
          header.style.width = `${{nextWidth}}px`;
          header.style.minWidth = `${{nextWidth}}px`;
        }};
        const onUp = () => {{
          document.removeEventListener("mousemove", onMove);
          document.removeEventListener("mouseup", onUp);
        }};
        handle.addEventListener("mousedown", (event) => {{
          startX = event.clientX;
          startWidth = header.getBoundingClientRect().width;
          document.addEventListener("mousemove", onMove);
          document.addEventListener("mouseup", onUp);
          event.preventDefault();
        }});
        header.appendChild(handle);
      }});
    }}

    function loadReviewState() {{
      try {{
        const raw = window.localStorage.getItem(reportStorageKey);
        const parsed = raw ? JSON.parse(raw) : {{}};
        return parsed && typeof parsed === "object" ? parsed : {{}};
      }} catch (_error) {{
        return {{}};
      }}
    }}

    function normalizeReviewEntry(value) {{
      if (typeof value === "string") {{
        return {{ checkedBy: value.trim(), comment: "" }};
      }}
      if (value && typeof value === "object") {{
        return {{
          checkedBy: typeof value.checkedBy === "string" ? value.checkedBy.trim() : "",
          comment: typeof value.comment === "string" ? value.comment : "",
        }};
      }}
      return {{ checkedBy: "", comment: "" }};
    }}

    function normalizeReviewState(value) {{
      if (!value || typeof value !== "object") {{
        return {{}};
      }}
      const normalized = {{}};
      Object.entries(value).forEach(([rowKey, entry]) => {{
        const reviewEntry = normalizeReviewEntry(entry);
        if (reviewEntry.checkedBy || reviewEntry.comment.trim()) {{
          normalized[rowKey] = reviewEntry;
        }}
      }});
      return normalized;
    }}

    function saveLocalReviewState() {{
      window.localStorage.setItem(reportStorageKey, JSON.stringify(reviewState));
    }}

    async function loadRemoteReviewState() {{
      if (window.location.protocol === "file:") {{
        return null;
      }}
      try {{
        const response = await window.fetch(reviewApiUrl, {{
          method: "GET",
          headers: {{ "Accept": "application/json" }},
        }});
        if (!response.ok) {{
          return null;
        }}
        const payload = await response.json();
        return normalizeReviewState(payload.rows);
      }} catch (_error) {{
        return null;
      }}
    }}

    async function saveReviewState() {{
      saveLocalReviewState();
      if (window.location.protocol === "file:") {{
        return;
      }}
      try {{
        await window.fetch(reviewApiUrl, {{
          method: "PUT",
          headers: {{ "Content-Type": "application/json" }},
          body: JSON.stringify({{ rows: reviewState }}),
        }});
      }} catch (_error) {{
        return;
      }}
    }}

    function showToast(message) {{
      if (!toast) return;
      toast.textContent = message;
      toast.classList.add("visible");
      if (toastTimer) {{
        window.clearTimeout(toastTimer);
      }}
      toastTimer = window.setTimeout(() => {{
        toast.classList.remove("visible");
      }}, 2400);
    }}

    async function copyText(value) {{
      if (navigator.clipboard && navigator.clipboard.writeText) {{
        await navigator.clipboard.writeText(value);
        return;
      }}
      const helper = document.createElement("textarea");
      helper.value = value;
      helper.setAttribute("readonly", "");
      helper.style.position = "absolute";
      helper.style.left = "-9999px";
      document.body.appendChild(helper);
      helper.select();
      document.execCommand("copy");
      document.body.removeChild(helper);
    }}

    function getStoredUsername() {{
      return (window.localStorage.getItem(usernameStorageKey) || "").trim();
    }}

    function ensureUsername() {{
      let username = getStoredUsername();
      if (username) {{
        return username;
      }}
      username = (window.prompt("Enter your username for Checked By") || "").trim();
      if (username) {{
        window.localStorage.setItem(usernameStorageKey, username);
      }}
      return username;
    }}

    function renderCheckedByCell(cell, checkedBy) {{
      const checkbox = cell.querySelector(".checked-by-checkbox");
      const name = cell.querySelector(".checked-by-name");
      const value = (checkedBy || "").trim();
      const row = cell.closest("tr.report-row");
      if (checkbox) {{
        checkbox.checked = Boolean(value);
      }}
      if (name) {{
        name.textContent = value || "-";
      }}
      if (row) {{
        row.dataset.checked = value ? "true" : "false";
      }}
    }}

    function renderCommentCell(cell, comment) {{
      const input = cell.querySelector(".comment-input");
      if (input) {{
        input.value = comment || "";
      }}
    }}

    function updateReviewState(rowKey, patch) {{
      const current = normalizeReviewEntry(reviewState[rowKey]);
      const next = {{
        checkedBy: Object.prototype.hasOwnProperty.call(patch, "checkedBy") ? patch.checkedBy : current.checkedBy,
        comment: Object.prototype.hasOwnProperty.call(patch, "comment") ? patch.comment : current.comment,
      }};
      if (!next.checkedBy && !next.comment.trim()) {{
        delete reviewState[rowKey];
      }} else {{
        reviewState[rowKey] = next;
      }}
      saveLocalReviewState();
      if (saveTimer) {{
        window.clearTimeout(saveTimer);
      }}
      saveTimer = window.setTimeout(() => {{
        void saveReviewState();
      }}, 250);
    }}

    function resetDetail(row) {{
      const detailId = row.dataset.detailTarget;
      if (!detailId) return;
      const detail = document.getElementById(detailId);
      if (detail) {{
        detail.hidden = true;
      }}
      const toggle = row.querySelector(".log-toggle");
      if (toggle) {{
        toggle.textContent = "View Source";
      }}
    }}

    function matches(row) {{
      const status = row.dataset.status || "";
      const categories = (row.dataset.categories || "").split("|").filter(Boolean);
      if (filters.group === "success" && status !== "success") {{
        return false;
      }}
      if (filters.group === "failed" && status === "success") {{
        return false;
      }}
      if (filters.category !== "all" && !categories.includes(filters.category)) {{
        return false;
      }}
      if (filters.status !== "all" && status !== filters.status) {{
        return false;
      }}
      if (filters.checked === "checked" && row.dataset.checked !== "true") {{
        return false;
      }}
      if (filters.checked === "unchecked" && row.dataset.checked === "true") {{
        return false;
      }}
      return true;
    }}

    function syncActiveStates() {{
      document.querySelectorAll(".filter-card, .filter-option").forEach((node) => {{
        const kind = node.dataset.filterKind;
        const value = node.dataset.filterValue;
        if (!kind || !value) return;
        node.classList.toggle("active", filters[kind] === value);
      }});
    }}

    function applyFilters() {{
      let visible = 0;
      reportRows.forEach((row) => {{
        const show = matches(row);
        row.hidden = !show;
        if (!show) {{
          resetDetail(row);
        }} else {{
          visible += 1;
        }}
      }});
      detailRows.forEach((row) => {{
        const parent = document.getElementById(row.dataset.parentRow);
        if (!parent || parent.hidden) {{
          row.hidden = true;
        }}
      }});
      if (filterState) {{
        filterState.textContent = `Showing ${{visible}} of {visible_rows} rows`;
      }}
      syncActiveStates();
    }}

    checkedByCells.forEach((cell) => {{
      const rowKey = cell.dataset.rowKey;
      const checkbox = cell.querySelector(".checked-by-checkbox");
      const initialCheckedBy = (cell.dataset.initialCheckedBy || "").trim();
      const savedEntry = normalizeReviewEntry(reviewState[rowKey]);
      const resolvedCheckedBy = savedEntry.checkedBy || initialCheckedBy;
      if (resolvedCheckedBy) {{
        reviewState[rowKey] = {{
          checkedBy: resolvedCheckedBy,
          comment: savedEntry.comment,
        }};
      }}
      renderCheckedByCell(cell, resolvedCheckedBy);
      if (!checkbox || !rowKey) {{
        return;
      }}
      checkbox.addEventListener("change", () => {{
        if (!checkbox.checked) {{
          renderCheckedByCell(cell, "");
          updateReviewState(rowKey, {{ checkedBy: "" }});
          return;
        }}
        const username = ensureUsername();
        if (!username) {{
          checkbox.checked = false;
          return;
        }}
        renderCheckedByCell(cell, username);
        updateReviewState(rowKey, {{ checkedBy: username }});
      }});
    }});
    commentCells.forEach((cell) => {{
      const rowKey = cell.dataset.rowKey;
      const initialComment = cell.dataset.initialComment || "";
      const savedEntry = normalizeReviewEntry(reviewState[rowKey]);
      const resolvedComment = savedEntry.comment || initialComment;
      if (rowKey && (savedEntry.checkedBy || resolvedComment)) {{
        reviewState[rowKey] = {{
          checkedBy: savedEntry.checkedBy,
          comment: resolvedComment,
        }};
      }}
      renderCommentCell(cell, resolvedComment);
      const input = cell.querySelector(".comment-input");
      if (!rowKey || !input) {{
        return;
      }}
      input.addEventListener("input", () => {{
        updateReviewState(rowKey, {{ comment: input.value }});
      }});
    }});
    saveLocalReviewState();

    async function hydrateRemoteReviewState() {{
      const remoteState = await loadRemoteReviewState();
      if (!remoteState) {{
        return;
      }}
      Object.entries(remoteState).forEach(([rowKey, entry]) => {{
        reviewState[rowKey] = normalizeReviewEntry(entry);
      }});
      checkedByCells.forEach((cell) => {{
        const rowKey = cell.dataset.rowKey;
        if (!rowKey) {{
          return;
        }}
        renderCheckedByCell(cell, normalizeReviewEntry(reviewState[rowKey]).checkedBy);
      }});
      commentCells.forEach((cell) => {{
        const rowKey = cell.dataset.rowKey;
        if (!rowKey) {{
          return;
        }}
        renderCommentCell(cell, normalizeReviewEntry(reviewState[rowKey]).comment);
      }});
      saveLocalReviewState();
      applyFilters();
    }}

    document.querySelectorAll(".filter-card, .filter-option").forEach((node) => {{
      node.addEventListener("click", () => {{
        const kind = node.dataset.filterKind;
        const value = node.dataset.filterValue;
        if (!kind || !value) return;
        if (kind === "group" && value === "all") {{
          filters.group = "all";
          filters.category = "all";
          filters.status = "all";
          filters.checked = "all";
        }} else {{
          filters[kind] = filters[kind] === value ? "all" : value;
        }}
        applyFilters();
      }});
    }});

    if (clearButton) {{
      clearButton.addEventListener("click", () => {{
        filters.group = "all";
        filters.category = "all";
        filters.status = "all";
        filters.checked = "all";
        applyFilters();
      }});
    }}

    document.querySelectorAll(".log-toggle").forEach((button) => {{
      button.addEventListener("click", () => {{
        const target = document.getElementById(button.dataset.target);
        if (!target) return;
        const parent = button.closest("tr.report-row");
        if (parent && parent.hidden) return;
        const hidden = target.hasAttribute("hidden");
        if (hidden) {{
          target.removeAttribute("hidden");
          button.textContent = "Hide Source";
        }} else {{
          target.setAttribute("hidden", "");
          button.textContent = "View Source";
        }}
      }});
    }});

    document.querySelectorAll(".log-open").forEach((button) => {{
      if (!canOpenLocalFiles) {{
        button.textContent = "Copy Path";
        button.title = "Browsers block opening local files from hosted reports. Click to copy the log path.";
      }}
      button.addEventListener("click", async () => {{
        const fileLink = button.dataset.fileLink || "";
        const logPath = button.dataset.logPath || "";
        if (canOpenLocalFiles && fileLink) {{
          window.open(fileLink, "_blank", "noopener,noreferrer");
          return;
        }}
        if (!logPath) {{
          showToast("Log path is not available for this row.");
          return;
        }}
        try {{
          await copyText(logPath);
          showToast("Local file path copied. Hosted reports cannot open local files directly.");
        }} catch (_error) {{
          showToast("Unable to copy the log path. Please use the path shown in the source panel.");
        }}
      }});
    }});

    enableColumnResize();
    applyFilters();
    void hydrateRemoteReviewState();
  }});
</script>
</head>
<body>
  <h1>{title}</h1>
  <div class="meta">Generated: {timestamp}</div>

  <div class="summary">
    <div class="summary-card filter-card active" data-filter-kind="group" data-filter-value="all">
      <h3>Total Models</h3>
      <div class="big">{total_models}</div>
    </div>
    <div class="summary-card filter-card" data-filter-kind="group" data-filter-value="success">
      <h3>Passed Models</h3>
      <div class="big" style="color: var(--success);">{passed_models}</div>
    </div>
    <div class="summary-card filter-card" data-filter-kind="group" data-filter-value="failed">
      <h3>Failed Models</h3>
      <div class="big" style="color: var(--danger);">{failed_models}</div>
    </div>
    <div class="summary-card">
      <h3>Errors by Category</h3>
      <table>{stats_rows}</table>
    </div>
    <div class="summary-card">
      <h3>Status</h3>
      <table>{status_rows}</table>
    </div>
    <div class="summary-card agent">
      <h3>Agent Assist</h3>
      <div class="agent-lines">
        <div><strong>Classifier</strong>: {classifier_model}</div>
        <div><strong>Debugger</strong>: {debugger_model}</div>
        <div><strong>AI-assisted fields</strong>: {assisted_fields}</div>
      </div>
    </div>
    <div class="summary-card source">
      <h3>Target Layout</h3>
      <div class="source-lines">
        <div><strong>Target Dir</strong>: {target_dir}</div>
        <div><strong>Layout Config</strong>: {layout_config}</div>
        <div><strong>Discovery</strong>: {discovery_rule}</div>
      </div>
      <div class="source-tree">{source_tree}</div>
    </div>
  </div>

  <div class="toolbar">
    <div class="toolbar-left">
      <div id="filter-state">Showing {visible_rows} of {visible_rows} rows</div>
      <div class="filter-pills" aria-label="Checked filter">
        <span class="filter-pills-label">Checked</span>
        <button type="button" class="filter-option active" data-filter-kind="checked" data-filter-value="all">All</button>
        <button type="button" class="filter-option" data-filter-kind="checked" data-filter-value="checked">Checked</button>
        <button type="button" class="filter-option" data-filter-kind="checked" data-filter-value="unchecked">Unchecked</button>
      </div>
    </div>
    <button type="button" id="clear-filters">Clear Filters</button>
  </div>

  <div class="table-wrapper">
    <table class="report">
      <thead><tr>{header_cells}</tr></thead>
      <tbody>
{body_rows}
      </tbody>
    </table>
  </div>

  <footer>model-test-agent v0.1.0</footer>
  <div id="toast" class="toast" aria-live="polite"></div>
</body>
</html>
"""
