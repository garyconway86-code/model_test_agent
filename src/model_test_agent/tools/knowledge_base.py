"""Minimal local-file RAG for MVP use."""

from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_SUPPORTED_SUFFIXES = {".txt", ".md", ".json", ".csv", ".log", ".xlsx"}
_TOKEN_RE = re.compile(r"[A-Za-z0-9_./:-]+")


@dataclass(frozen=True)
class KnowledgeChunk:
    source: str
    title: str
    text: str


class KnowledgeBase:
    """Read a small local knowledge directory and return relevant chunks."""

    def __init__(
        self,
        root_dir: str | Path,
        chunk_size: int = 1200,
        overlap: int = 120,
        excerpt_chars: int = 900,
        embedding_input_chars: int = 4000,
    ) -> None:
        self.root_dir = Path(root_dir)
        self.chunk_size = chunk_size
        self.overlap = overlap
        self.excerpt_chars = excerpt_chars
        self.embedding_input_chars = embedding_input_chars
        self._chunks: list[KnowledgeChunk] | None = None
        self._embeddings: list[list[float]] | None = None

    def exists(self) -> bool:
        return self.root_dir.exists() and self.root_dir.is_dir()

    def search(
        self,
        query: str,
        top_k: int = 3,
        embedding_client: Any = None,
        embedding_model: str = "",
    ) -> list[dict[str, Any]]:
        chunks = self._load_chunks()
        if not chunks or not query.strip():
            return []

        lexical_scores = self._lexical_scores(query, chunks)
        if embedding_client and embedding_model:
            try:
                semantic_scores = self._semantic_scores(query, chunks, embedding_client, embedding_model)
            except Exception:
                semantic_scores = {}
        else:
            semantic_scores = {}

        ranked: list[tuple[float, KnowledgeChunk]] = []
        for index, chunk in enumerate(chunks):
            score = lexical_scores.get(index, 0.0) + semantic_scores.get(index, 0.0)
            if score > 0:
                ranked.append((score, chunk))

        ranked.sort(key=lambda item: item[0], reverse=True)
        return [
            {
                "id": f"doc-{idx}",
                "kind": "document",
                "source": chunk.source,
                "title": chunk.title,
                "excerpt": chunk.text[: self.excerpt_chars],
                "score": round(score, 3),
            }
            for idx, (score, chunk) in enumerate(ranked[:top_k], start=1)
        ]

    def _load_chunks(self) -> list[KnowledgeChunk]:
        if self._chunks is not None:
            return self._chunks
        if not self.exists():
            self._chunks = []
            return self._chunks

        chunks: list[KnowledgeChunk] = []
        for path in sorted(self.root_dir.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in _SUPPORTED_SUFFIXES:
                continue
            content = self._read_file(path)
            if not content.strip():
                continue
            relative = str(path.relative_to(self.root_dir))
            for piece in self._chunk_text(content):
                chunks.append(KnowledgeChunk(source=relative, title=path.name, text=piece))
        self._chunks = chunks
        return chunks

    def _read_file(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".log"}:
            return path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".json":
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
            except json.JSONDecodeError:
                return path.read_text(encoding="utf-8", errors="replace")
            return json.dumps(data, ensure_ascii=False, indent=2)
        if suffix == ".csv":
            rows: list[str] = []
            with open(path, encoding="utf-8", errors="replace", newline="") as fh:
                reader = csv.reader(fh)
                for row in reader:
                    rows.append(" | ".join(cell.strip() for cell in row))
            return "\n".join(rows)
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheets: list[str] = []
            for sheet in workbook.worksheets:
                rows = [f"[sheet] {sheet.title}"]
                for row in sheet.iter_rows(values_only=True):
                    values = [str(cell).strip() for cell in row if cell not in (None, "")]
                    if values:
                        rows.append(" | ".join(values))
                sheets.append("\n".join(rows))
            return "\n\n".join(sheets)
        return ""

    def _chunk_text(self, text: str) -> list[str]:
        normalized = text.replace("\r\n", "\n")
        if len(normalized) <= self.chunk_size:
            return [normalized]

        chunks: list[str] = []
        start = 0
        while start < len(normalized):
            end = min(len(normalized), start + self.chunk_size)
            chunks.append(normalized[start:end])
            if end >= len(normalized):
                break
            start = max(end - self.overlap, start + 1)
        return chunks

    def _lexical_scores(self, query: str, chunks: list[KnowledgeChunk]) -> dict[int, float]:
        query_tokens = self._tokenize(query)
        scores: dict[int, float] = {}
        if not query_tokens:
            return scores
        for index, chunk in enumerate(chunks):
            overlap = query_tokens & self._tokenize(chunk.text)
            if overlap:
                scores[index] = float(len(overlap))
        return scores

    def _semantic_scores(
        self,
        query: str,
        chunks: list[KnowledgeChunk],
        embedding_client: Any,
        embedding_model: str,
    ) -> dict[int, float]:
        query_embedding = embedding_client.embeddings.create(model=embedding_model, input=[query]).data[0].embedding
        chunk_embeddings = self._load_embeddings(chunks, embedding_client, embedding_model)
        scores: dict[int, float] = {}
        for index, embedding in enumerate(chunk_embeddings):
            similarity = self._cosine_similarity(query_embedding, embedding)
            if similarity > 0:
                scores[index] = similarity * 3.0
        return scores

    def _load_embeddings(
        self,
        chunks: list[KnowledgeChunk],
        embedding_client: Any,
        embedding_model: str,
    ) -> list[list[float]]:
        if self._embeddings is not None:
            return self._embeddings
        inputs = [chunk.text[: self.embedding_input_chars] for chunk in chunks]
        data = embedding_client.embeddings.create(model=embedding_model, input=inputs).data
        self._embeddings = [item.embedding for item in data]
        return self._embeddings

    @staticmethod
    def _tokenize(text: str) -> set[str]:
        return {token.lower() for token in _TOKEN_RE.findall(text)}

    @staticmethod
    def _cosine_similarity(left: list[float], right: list[float]) -> float:
        numerator = sum(a * b for a, b in zip(left, right))
        left_norm = math.sqrt(sum(a * a for a in left))
        right_norm = math.sqrt(sum(b * b for b in right))
        if left_norm == 0 or right_norm == 0:
            return 0.0
        return numerator / (left_norm * right_norm)
